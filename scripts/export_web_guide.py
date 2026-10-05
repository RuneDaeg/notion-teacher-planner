#!/usr/bin/env python3
"""Export the reviewed Word guide to hosted and self-contained HTML.

Run with a Python environment containing python-docx and pypdf. The DOCX and
reviewed PDF are the source artifacts; this script does not rewrite either.
"""

from __future__ import annotations

import argparse
from html import escape
from pathlib import Path
import re
import sys
from urllib.parse import urlsplit

from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
ORIGIN = "https://notion-teacher-planner.notion-teacher-planner-cloudflare.workers.dev"
REPOSITORY = "https://github.com/RuneDaeg/notion-teacher-planner"
GUIDE_NAME = "교무수첩_사용안내서"


def safe_url(value: str) -> str:
    """DOCX links are data, never executable markup or script URLs."""
    parsed = urlsplit(value)
    if parsed.scheme not in {"https", "http", "mailto"}:
        raise ValueError(f"Unsupported guide hyperlink scheme: {parsed.scheme!r}")
    return escape(value, quote=True)


def render_run(run: Run) -> str:
    value = escape(run.text).replace("\n", "<br>\n").replace("\t", "&#9;")
    if run.bold:
        value = f"<strong>{value}</strong>"
    if run.italic:
        value = f"<em>{value}</em>"
    return value


def inline(paragraph: Paragraph) -> str:
    pieces = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            label = "".join(render_run(run) for run in item.runs)
            if item.url:
                pieces.append(f'<a href="{safe_url(item.url)}">{label}</a>')
            else:
                pieces.append(label)
        else:
            pieces.append(render_run(item))
    return "".join(pieces)


def render_table(table: Table, label: str) -> str:
    rows = []
    for row_number, row in enumerate(table.rows):
        tag = "th" if row_number == 0 else "td"
        attribute = ' scope="col"' if row_number == 0 else ""
        cells = []
        for cell in row.cells:
            content = "".join(f"<p>{inline(p)}</p>" for p in cell.paragraphs if p.text.strip())
            cells.append(f"<{tag}{attribute}>{content}</{tag}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    if not rows:
        return ""
    return (
        f'<div class="table-scroll" tabindex="0" role="region" aria-label="{escape(label, quote=True)} 표">'
        f"<table><thead>{rows[0]}</thead><tbody>{''.join(rows[1:])}</tbody></table></div>"
    )


def render_body(document) -> tuple[str, list[tuple[str, str]]]:
    """Preserve body order, including tables, and give each chapter an anchor."""
    parts = ['<section class="guide-section introduction" id="start" aria-label="시작 안내">']
    toc = [("start", "시작 안내")]
    section_number = 0
    list_kind = None
    preceding_heading = "사용 안내"

    def end_list():
        nonlocal list_kind
        if list_kind:
            parts.append(f"</{list_kind}>")
            list_kind = None

    for item in document.iter_inner_content():
        if isinstance(item, Table):
            end_list()
            parts.append(render_table(item, preceding_heading))
            continue
        text = item.text.strip()
        if not text:
            continue
        style = item.style.name
        if style == "Title":
            end_list()
            parts.append(f'<h1 id="guide-title">{inline(item)}</h1>')
            continue
        if style == "Subtitle":
            end_list()
            parts.append(f'<p class="subtitle">{inline(item)}</p>')
            continue
        if "문서 버전" in text and section_number == 0:
            end_list()
            parts.append(f'<p class="document-meta">{inline(item)}</p>')
            continue
        if style == "Heading 1":
            end_list()
            section_number += 1
            anchor = f"chapter-{section_number}"
            toc.append((anchor, text))
            parts.append(f'</section><section class="guide-section" id="{anchor}" aria-labelledby="{anchor}-title">')
            parts.append(f'<h2 id="{anchor}-title">{inline(item)}</h2>')
            preceding_heading = text
            continue
        if style.startswith("Heading"):
            end_list()
            level = "h3" if section_number else "h2"
            parts.append(f"<{level}>{inline(item)}</{level}>")
            preceding_heading = text
            continue

        numbered = re.match(r"^\d+\.\s+", text)
        kind = "ul" if style.startswith("List Bullet") else "ol" if numbered or style.startswith("List Number") else None
        if kind:
            if list_kind != kind:
                end_list()
                # The source already writes numbered prefixes; retain them exactly.
                marker_class = ' class="written-steps"' if numbered else ""
                parts.append(f"<{kind}{marker_class}>")
                list_kind = kind
            parts.append(f"<li>{inline(item)}</li>")
        else:
            end_list()
            if text.startswith(("“", "AI 요청 예")):
                parts.append(f'<blockquote class="request-example"><p>{inline(item)}</p></blockquote>')
            else:
                parts.append(f"<p>{inline(item)}</p>")
    end_list()
    parts.append("</section>")
    return "\n".join(parts), toc


def render_document(document, css: str, *, portable: bool, version: str) -> str:
    body, contents = render_body(document)
    base = ORIGIN if portable else ""
    style = f"<style>\n{css}\n</style>" if portable else '<link rel="stylesheet" href="/guide.css">'
    toc = "\n".join(f'<li><a href="#{anchor}">{escape(title)}</a></li>' for anchor, title in contents)
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="guide-version" content="{escape(version, quote=True)}">
  <meta name="description" content="Notion 교무수첩의 설치, 학생 상담, 시간표, 중식·학사일정부터 드래그 배치 미리보기와 AI 수정 요청까지 읽는 교사용 안내서.">
  <title>Notion 교무수첩 사용 안내서</title>
  {style}
</head>
<body>
  <a class="skip-link" href="#guide-title">안내서 본문으로 건너뛰기</a>
  <header class="site-header">
    <a class="brand" href="{base}/setup"><span class="brand-mark" aria-hidden="true">▤</span> Notion 교무수첩</a>
    <nav aria-label="사이트 메뉴"><a href="{base}/setup">내 수첩 준비하기</a><a href="{base}/preview">전체 배치 미리보기</a><a href="{base}/edit">내 수첩 수정하기</a><a href="{REPOSITORY}">GitHub</a></nav>
  </header>
  <div class="page-shell">
    <aside class="contents-rail">
      <details open>
        <summary>이 안내서의 목차</summary>
        <nav aria-label="안내서 목차"><ol>{toc}</ol></nav>
      </details>
      <p class="rail-note">필요한 항목부터 읽어도 좋아요.<br>파일로 받아 두면 오프라인에서도 볼 수 있습니다.</p>
    </aside>
    <main id="main-content">
      <div class="reading-tools" aria-label="다른 형식으로 읽기">
        <span>사용 안내서</span>
        <div class="file-actions">
          <a href="{base}/teacher-planner-guide.pdf" target="_blank" rel="noopener">PDF 보기 <span class="sr-only">(새 탭)</span><span aria-hidden="true">↗</span></a>
          <a href="{base}/teacher-planner-guide.pdf" download="{GUIDE_NAME}.pdf">PDF 받기 <span aria-hidden="true">↓</span></a>
          <a href="{base}/teacher-planner-guide.html" download="{GUIDE_NAME}.html">HTML 받기 <span aria-hidden="true">↓</span></a>
        </div>
      </div>
      <p class="format-note">브라우저 인쇄(Ctrl/Cmd+P)로 인쇄할 수도 있습니다. 본문의 쪽수는 PDF 기준입니다.</p>
      <article class="guide-document" aria-labelledby="guide-title">
{body}
      </article>
      <footer class="document-footer">
        <p>내 학교와 담당 수업에 맞는 수첩을 준비하세요.</p>
        <a class="start-link" href="{base}/setup">웹 질문지로 시작하기 <span aria-hidden="true">→</span></a>
        <a class="back-top" href="#guide-title">맨 위로</a>
        <p class="credits"><span>made by 여광재(온양고)</span><span>made with <a href="https://dorms.school">DoRms</a></span></p>
      </footer>
    </main>
  </div>
</body>
</html>
"""


def source_version(document, pdf_path: Path) -> str:
    docx_text = "\n".join(p.text for p in document.paragraphs)
    match = re.search(r"(\d{4}년 \d{1,2}월 \d{1,2}일) 기준\s*·\s*문서 버전 ([\d.]+)", docx_text)
    if not match:
        raise ValueError("DOCX is missing its guide date and version.")
    pdf_text = " ".join((p.extract_text() or "") for p in PdfReader(pdf_path).pages)
    normalized = re.sub(r"\s+", " ", pdf_text)
    if match.group(1) not in normalized or f"문서 버전 {match.group(2)}" not in normalized:
        raise ValueError("DOCX and PDF date/version differ. Regenerate and review the PDF first.")
    return match.group(2)


def outputs(root: Path = ROOT) -> tuple[dict[Path, bytes], str]:
    docx_path = root / "output" / "docx" / f"{GUIDE_NAME}.docx"
    pdf_path = root / "output" / "pdf" / f"{GUIDE_NAME}.pdf"
    css_path = root / "cloud" / "web" / "guide.css"
    for path in (docx_path, pdf_path, css_path):
        if not path.is_file():
            raise FileNotFoundError(f"Required source is missing: {path.relative_to(root)}")
    document = Document(docx_path)
    version = source_version(document, pdf_path)
    css = css_path.read_text(encoding="utf-8")
    portable = render_document(document, css, portable=True, version=version).encode("utf-8")
    return {
        root / "cloud" / "web" / "guide.html": render_document(document, css, portable=False, version=version).encode("utf-8"),
        root / "output" / "html" / f"{GUIDE_NAME}.html": portable,
        root / "cloud" / "web" / "teacher-planner-guide.html": portable,
        root / "cloud" / "web" / "teacher-planner-guide.pdf": pdf_path.read_bytes(),
    }, version


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Check exports match the reviewed source without writing files.")
    args = parser.parse_args()
    try:
        files, version = outputs()
        mismatches = []
        for path, data in files.items():
            if args.check:
                if not path.is_file() or path.read_bytes() != data:
                    mismatches.append(str(path.relative_to(ROOT)))
            else:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(data)
        if mismatches:
            print("Out-of-date guide exports: " + ", ".join(mismatches), file=sys.stderr)
            return 1
        print(f"Guide v{version}: {len(files)} exports {'verified' if args.check else 'written'}.")
        return 0
    except (FileNotFoundError, ValueError) as error:
        print(str(error), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
